import pandas as pd
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression

# =========================
# LER DOCUMENTO WORD
# =========================

doc = Document("dataset.docx")

print("Parágrafos:", len(doc.paragraphs))
print("Tabelas:", len(doc.tables))

headlines = []
labels = []

# =========================
# LER TABELAS
# =========================

for table in doc.tables:

    for i, row in enumerate(table.rows):

        cells = [cell.text.strip() for cell in row.cells]

        # Ignorar cabeçalho
        if i == 0:
            continue

        # Verificar se existem pelo menos 5 colunas
        if len(cells) >= 5:

            headline = cells[1]
            label = cells[4]

            if label in ["Fake", "Real"]:

                headlines.append(headline)
                labels.append(label)

# =========================
# CRIAR DATAFRAME
# =========================

df = pd.DataFrame({
    "Headline": headlines,
    "Label": labels
})

print("\nNúmero de notícias carregadas:", len(df))
print("\nPrimeiras linhas:")
print(df.head())

if len(df) == 0:
    print("\nERRO: Não foi encontrada nenhuma notícia.")
    exit()

# =========================
# CONVERTER LABELS
# =========================

y = df["Label"].map({
    "Fake": 1,
    "Real": 0
})

# =========================
# TF-IDF
# =========================

vectorizer = TfidfVectorizer()

X = vectorizer.fit_transform(df["Headline"])

# =========================
# TREINAR MODELO
# =========================

model = LogisticRegression(max_iter=1000)

model.fit(X, y)

print("\nModelo treinado com sucesso!")

# =========================
# TESTES
# =========================

test_data = [

    # Casos fáceis
    ("T01", "Garlic tea completely prevents heart attacks forever", 1),
    ("T02", "Doctors recommend regular exercise for heart health", 0),
    ("T03", "Orange soda cleans blood vessels instantly", 1),
    ("T04", "Researchers improve wearable heart monitoring technology", 0),

    # Casos médios
    ("T05", "Scientists discover new treatment for heart patients", 0),
    ("T06", "Honey cures blocked arteries overnight", 1),
    ("T07", "Cardiologists encourage healthy eating habits", 0),
    ("T08", "Magnets improve blood circulation instantly", 1),

    # Casos difíceis
    ("T09", "New heart medication shows promising early results", 0),
    ("T10", "Daily lemon juice may reduce heart attack risk", 1),
    ("T11", "Experts study effects of sleep on heart health", 0),
    ("T12", "Special tea completely eliminates cholesterol", 1),

    # Casos ambíguos
    ("T13", "Researchers investigate garlic compounds for cardiovascular health", 0),
    ("T14", "Garlic supplements guarantee protection from heart attacks", 1),
    ("T15", "Scientists analyse diet and cardiovascular disease", 0),
    ("T16", "Natural remedy repairs damaged arteries instantly", 1),

    # Casos extra
    ("T17", "Hospitals launch new emergency cardiac response system", 0),
    ("T18", "Walking five minutes per week prevents all heart disease", 1),
    ("T19", "New study links smoking cessation to lower cardiac risk", 0),
    ("T20", "Chocolate milk permanently removes artery blockage", 1)
]

# Preparar dados de teste
sentences = [headline for _, headline, _ in test_data]

X_test = vectorizer.transform(sentences)

# Previsões
predictions = model.predict(X_test)

# Probabilidades
probabilities = model.predict_proba(X_test)

# =========================
# RESULTADOS
# =========================

print("\nRESULTADOS\n")

acertos = 0

for i, ((news_id, headline, esperado), pred) in enumerate(zip(test_data, predictions)):

    previsto = "Fake News" if pred == 1 else "Real News"
    real = "Fake News" if esperado == 1 else "Real News"

    confidence = max(probabilities[i]) * 100

    correto = pred == esperado

    if correto:
        acertos += 1

    print(f"ID: {news_id}")
    print(f"Headline: {headline}")
    print(f"Valor Real: {real}")
    print(f"Previsão: {previsto}")
    print(f"Confiança: {confidence:.2f}%")

    if correto:
        print("✓ ACERTOU")
    else:
        print("✗ ERROU")

    print("-" * 60)

# =========================
# ESTATÍSTICAS
# =========================

total = len(test_data)

precisao = (acertos / total) * 100

print("\nRESUMO FINAL")
print(f"Total de testes: {total}")
print(f"Acertos: {acertos}")
print(f"Erros: {total - acertos}")
print(f"Precisão: {precisao:.2f}%")

# =========================
# TESTE LIVRE
# =========================

while True:

    frase = input("\nIntroduza uma headline (ou 'sair'): ")

    if frase.lower() == "sair":
        print("Programa terminado.")
        break

    X_new = vectorizer.transform([frase])

    prediction = model.predict(X_new)
    probability = model.predict_proba(X_new)

    confidence = max(probability[0]) * 100

    if prediction[0] == 1:
        print(f"Resultado: Fake News ({confidence:.2f}% confiança)")
    else:
        print(f"Resultado: Real News ({confidence:.2f}% confiança)")
